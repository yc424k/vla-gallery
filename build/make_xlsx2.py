import json, glob, re, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

P = sys.argv[1]
OUT = sys.argv[2]
FONT = "맑은 고딕"

KEYS = ["모델","구분","유형","도메인","발표","기관","논문","논문 제목","논문 링크","관련 논문","관련 논문 링크","가중치 공개","Hugging Face","아키텍처","백본 LLM","LLM 계열","LLM 크기(B)","비전 인코더","액션 헤드","액션 헤드 종류","액션 청크","제어 주기","베이스 종류","베이스 제어 방식","베이스 제어 설명","백본·액션 표현","핵심 특징","파라미터","파라미터(B)","사전학습 가속기","사전학습 기간","사전학습 스텝·배치","사전학습 GPU·TPU-시간","시간 근거·단위","파인튜닝 가속기·시간","강화학습","강화학습 방식","강화학습 사용 위치","강화학습 근거 URL","추론 장치·속도","추론 메모리(보고)","가중치 bf16(계산, GB)","실행 위치","태스크·벤치마크","로봇·엠바디먼트","데이터 종류","데이터 규모","데이터 시간(h)","주요 데이터셋","비고","검증","1차 출처 URL"]
DKEYS = ["데이터셋","발표","기관","종류","에피소드·궤적","시간(h)","프레임","태스크·스킬","엠바디먼트","접근·라이선스","사용 모델","비고","검증","출처 URL"]
ORDER = ["참고 기준(2024 이전)","VA","자기회귀·잠재 액션 VLA","π 계열","Flow·Diffusion VLA","산업계 모델","휴머노이드 VLA","휴머노이드 전신 제어기","모바일 매니퓰레이션","내비게이션","2026 신규"]

def load(pattern):
    rows=[]
    for f in sorted(glob.glob(f"{P}/{pattern}")):
        for l in open(f, encoding="utf-8"):
            if l.strip(): rows.append(json.loads(l))
    return rows

models = load("rows_w[1-6].jsonl") + load("rows_w8.jsonl")
GD = {g["name"]: g for g in json.load(open(sys.argv[3], encoding="utf-8"))}
for r in models:
    g = GD[r["모델"]]
    r["백본 LLM"] = g["llm"]; r["LLM 계열"] = g["fam"]; r["LLM 크기(B)"] = g["lb"]
    r["비전 인코더"] = g["vis"] or "—"; r["액션 헤드"] = g["head"]; r["액션 헤드 종류"] = g["hk"]
    r["액션 청크"] = g["chunk"]; r["제어 주기"] = g["hz"]
    r["추론 장치·속도"] = r.pop("추론"); r["추론 메모리(보고)"] = g["mem"] or "미공개"
    r["가중치 bf16(계산, GB)"] = g["w16"]; r["실행 위치"] = g["dep"] or "미공개"
    r["강화학습"] = g["rl"]; r["강화학습 방식"] = g["rlcat"]; r["강화학습 사용 위치"] = g["rlwhere"]; r["강화학습 근거 URL"] = g["rlsrc"] or ""
    pp, pr = g["paper"], g["prel"]
    r["논문"] = {"논문": "있음", "보고서": "기술 보고서", "없음": "없음"}[pp["k"]]
    r["논문 제목"] = (pp["t"] + (f" (arXiv {pp['id']})" if pp["id"] else "") + (f" — {pp['n']}" if pp.get("n") else "")) if pp["u"] else pp["n"]
    r["논문 링크"] = pp["u"] or ""
    r["Hugging Face"] = "\n".join(g["hf"])
    r["관련 논문"] = (pr["t"] + (f" (arXiv {pr['id']})" if pr["id"] else "") + " — " + pr["why"]) if pr else ""
    r["관련 논문 링크"] = pr["u"] if pr else ""
    bb = g["base"]; mob = [t for t in bb["types"] if t != "고정"]
    r["베이스 종류"] = " · ".join(bb["types"]); r["베이스 제어 방식"] = bb.get("cat") or ("베이스 제어 없음" if not mob else "미분류"); r["베이스 제어 설명"] = bb["how"]

datasets = load("rows_w7.jsonl")
assert all(len(r)==52 for r in models), [len(r) for r in models][:3] and all(len(r)==14 for r in datasets)
models.sort(key=lambda r: (ORDER.index(r["구분"]) if r["구분"] in ORDER else 99, str(r["발표"] or "")))
datasets.sort(key=lambda r: str(r["발표"] or ""))

hdr_font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
hdr_fill = PatternFill("solid", fgColor="1F4E79")
body_font = Font(name=FONT, size=9)
link_font = Font(name=FONT, size=9, color="0563C1", underline="single")
thin = Side(style="thin", color="D0D7DE")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap_top = Alignment(wrap_text=True, vertical="top")
center_top = Alignment(horizontal="center", vertical="top", wrap_text=True)
band_fill = PatternFill("solid", fgColor="F3F6FA")

def write_sheet(ws, keys, rows, widths, numeric, urlkey, numfmt):
    ws.append(keys)
    for c in ws[1]:
        c.font = hdr_font; c.fill = hdr_fill; c.alignment = center_top; c.border = border
    for r in rows:
        ws.append([r.get(k) for k in keys])
    for i, row in enumerate(ws.iter_rows(min_row=2, max_row=ws.max_row), start=2):
        prev = ws.cell(row=i-1, column=2).value if i>2 else None
        for c in row:
            k = keys[c.column-1]
            c.border = border
            if k in (urlkey if isinstance(urlkey, set) else {urlkey}) and isinstance(c.value, str) and c.value.startswith("http"):
                c.hyperlink = c.value.split()[0]
                c.font = link_font
            else:
                c.font = body_font
            if k in numeric:
                c.alignment = Alignment(horizontal="right", vertical="top")
                c.number_format = numfmt.get(k, "#,##0")
            else:
                c.alignment = wrap_top
    for idx, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = w
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(keys))}{ws.max_row}"
    ws.row_dimensions[1].height = 30

wb = Workbook()
ws = wb.active; ws.title = "모델 비교"
mwidths = [22,16,8,12,9,22,10,46,34,46,34,9,40,13,26,10,9,24,28,13,14,14,16,22,44,28,40,18,9,22,18,20,12,26,30,9,16,40,28,30,24,11,16,34,22,22,32,10,26,34,10,40]
write_sheet(ws, KEYS, models, mwidths, {"파라미터(B)","LLM 크기(B)","사전학습 GPU·TPU-시간","가중치 bf16(계산, GB)","데이터 시간(h)"}, {"1차 출처 URL","논문 링크","관련 논문 링크","강화학습 근거 URL","Hugging Face"},
            {"파라미터(B)":"0.####","LLM 크기(B)":"0.##","사전학습 GPU·TPU-시간":"#,##0.##","가중치 bf16(계산, GB)":"#,##0.##","데이터 시간(h)":"#,##0.#"})
# band rows by 구분 to make groups scannable
cur=None; shade=False
for i in range(2, ws.max_row+1):
    g = ws.cell(row=i, column=2).value
    if g != cur: cur=g; shade = not shade
    if shade:
        for j in range(1, len(KEYS)+1):
            ws.cell(row=i, column=j).fill = band_fill

ws2 = wb.create_sheet("데이터셋")
dwidths = [26,9,24,22,26,10,18,22,24,22,34,34,10,40]
write_sheet(ws2, DKEYS, datasets, dwidths, {"시간(h)"}, "출처 URL", {"시간(h)":"#,##0.#"})

# 출처 sheet
ws3 = wb.create_sheet("출처")
labels = {"w1":"참고 기준·VA","w2":"자기회귀·잠재 액션 VLA","w3":"π 계열·Flow/Diffusion VLA","w4":"산업계 모델","w5":"휴머노이드","w6":"2026 신규","w8":"모바일 매니퓰레이션·내비게이션","w7":"데이터셋"}
srows=[]
link_re = re.compile(r"\[([^\]]+)\]\((https?://[^)\s]+)\)")
for key,label in labels.items():
    f = f"{P}/sources_{key}.md"
    try:
        txt = open(f, encoding="utf-8").read()
    except FileNotFoundError:
        txt = ""
    for t,u in link_re.findall(txt):
        srows.append({"구분":label,"제목":t,"URL":u})
# primary sources of every model/dataset row too
for r in models:
    u = r.get("1차 출처 URL")
    if isinstance(u,str) and u.startswith("http"):
        srows.append({"구분":"모델 1차 출처 · "+r["구분"],"제목":r["모델"],"URL":u.split()[0]})
for r in datasets:
    u = r.get("출처 URL")
    if isinstance(u,str) and u.startswith("http"):
        srows.append({"구분":"데이터셋 1차 출처","제목":r["데이터셋"],"URL":u.split()[0]})
seen=set(); uniq=[]
for s in srows:
    k=(s["구분"],s["URL"])
    if k in seen: continue
    seen.add(k); uniq.append(s)
write_sheet(ws3, ["구분","제목","URL"], uniq, [30,60,70], set(), "URL", {})

# 표기 기준 sheet
ws4 = wb.create_sheet("표기 기준")
notes = [
 ("기준일","2026-10-14 (모바일 매니퓰레이션·내비게이션 모델과 베이스 제어 열 추가. 기존 모델 수치는 2026-10-06 기준)"),
 ("범위","2024년 이후 발표된 로봇 제어용 VA·VLA 모델 위주(2024년 이전 대표 모델은 '참고 기준'), 트랜스포머 기반 위주. 매니퓰레이션·휴머노이드 전신 제어·모바일 베이스 제어(모바일 매니퓰레이션·내비게이션, 2026-10-14 추가). 모바일 매니퓰레이션·내비게이션 분류에는 GNM·ViNT·NoMaD·HarmonicMM·QUAR-VLA 같은 2024년 이전 기준 모델도 함께 넣음. 자율주행 차량 전용 모델 제외."),
 ("VA / VLA","VA(Vision-Action): 언어 입력 없이 이미지·로봇 상태에서 액션을 내는 정책. VLA(Vision-Language-Action): 언어 지시를 함께 받는 정책. '기타'는 world action model 등."),
 ("아키텍처","'비트랜스포머(비교용)'은 Diffusion Policy(U-Net), MLP 전신 제어기처럼 영향력이 커서 비교용으로만 넣은 모델."),
 ("발표","첫 arXiv 버전과 첫 공식 발표 중 이른 쪽의 연·월(YYYY-MM)."),
 ("미공개","논문·README·모델 카드·공식 블로그 어디에도 값이 없음. 추정하지 않음."),
 ("(2차)","1차 출처가 아닌 기사·타사 논문·커뮤니티 자료에서 가져온 값."),
 ("†","원문 대조로 확인하지 못한 값."),
 ("A/B 표기","출처끼리 값이 다를 때 둘 다 표기. 예: 7B(논문)/7.5B(README)."),
 ("사전학습 GPU·TPU-시간","출처가 밝힌 GPU-시간은 그대로(명시). 없으면 장수와 기간이 모두 숫자로 공개된 경우에만 장수 × 시간으로 계산(환산). '약 1개월'은 30일로 계산(환산(약)). 환산값은 가동률 100%를 가정한 상한이며, GPU 종류가 달라 다른 모델과 직접 비교할 수 없음. TPU와 GPU는 합산하지 않음."),
 ("검증","확인: 1차 출처 문장과 대조함 / 수정 반영: 대조 결과 바로잡음 / 확인불가: 원문에 접근하지 못함 / 미검증: 대조 대상 밖."),
 ("파라미터(B)","정렬용 숫자(십억 단위). 크기가 여러 개면 가장 큰 값이며, 비고에 설명."),
 ("백본 LLM","VLA 안의 언어 모델. '· PaliGemma 3B 구성'처럼 적힌 값은 논문이 VLM 이름만 밝혀, 해당 VLM의 공개 구성에서 LLM을 옮긴 것. LLM이 없는 모델은 '없음', 밝히지 않은 모델은 '미공개'."),
 ("가중치 bf16(계산, GB)","총 파라미터(B) × 2바이트 계산값. 활성값·KV 캐시·이미지 토큰을 포함하지 않으므로 실제 추론 메모리는 더 크다. 보고값이 있는 모델에서는 0~4 GB 더 컸다."),
 ("추론 메모리(보고)","논문·README·모델 카드가 밝힌 추론 시 GPU 메모리. 정밀도와 출처를 함께 적었다."),
 ("논문","있음 = 그 모델을 처음 소개한 논문(arXiv 초록 페이지 링크, 제목은 원문 그대로). 한 논문이 두 모델을 함께 소개하면 같은 논문을 적음. 기술 보고서 = arXiv 논문 없이 개발사가 낸 기술 보고서. 없음 = 블로그·모델 카드·저장소로만 공개. 관련 논문 = 자체 논문이 없을 때 공식 페이지가 밝힌 전작·기반 논문(근거 함께 표기). 2026-10-09 확인."),
 ("Hugging Face","모델 저자·개발사가 Hugging Face에 올린 공식 가중치 저장소(여러 개면 줄바꿈). 제3자 파인튜닝·재업로드·양자화본과 데이터셋만 있는 저장소는 제외. 2026-10-17 확인."),
 ("강화학습","사용 = 모델 자체를 RL로 학습·개선. 일부 = 하위 제어기·RL 생성 데이터·보상·선호 최적화(DPO)·가치 예측 등 일부에만 사용. 미사용 = 공개된 학습 과정이 모방·지도학습뿐. 미공개 = 학습 방식을 밝히지 않음. 근거 URL은 논문·공식 블로그."),
 ("베이스 종류","고정 = 바퀴·다리가 있어도 이 모델이 베이스를 움직이지 않음 / 바퀴형 / 사족보행 / 휴머노이드 보행 / 비행 / 시뮬 전용 = 시뮬레이션에서만 베이스를 움직임. 여러 종류를 움직이면 모두 표기."),
 ("베이스 제어 방식","베이스 속도·자세 직접 출력 / 웨이포인트 → 추종 제어기(PD·MPC·pure pursuit) / 이산 이동 행동(전진·회전·정지) / 명령·잠재 → RL 보행 제어기 / RL 정책이 직접 구동 / 전신 관절·자세 직접 출력 / 계획·최적화로 이동 결정 / 형식 미기재. 근거는 논문·공식 문서의 행동 공간 서술."),
 ("갤러리","구조도·필터·추론 메모리 계산기·두 모델 비교는 함께 공유한 '로봇 VLA 아키텍처 갤러리' 페이지에 있다."),
]
ws4.append(["항목","설명"])
for n in notes: ws4.append(list(n))
for c in ws4[1]:
    c.font = hdr_font; c.fill = hdr_fill; c.alignment = center_top; c.border = border
for row in ws4.iter_rows(min_row=2):
    for c in row:
        c.font = Font(name=FONT, size=10, bold=(c.column==1)); c.alignment = wrap_top; c.border = border
ws4.column_dimensions["A"].width = 22; ws4.column_dimensions["B"].width = 110

wb.move_sheet("표기 기준", offset=-3)
wb.active = 1
wb.save(OUT)
print("models", len(models), "datasets", len(datasets), "sources", len(uniq))
